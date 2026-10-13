"""Offline display/context regression; no ERP or settings database access."""
from __future__ import annotations

from io import BytesIO
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
import socket
import sys

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.services.product_flow_service import _prepare_display_df
from app.services.dashboard_lite_facts import _attach_inventory_status_and_frequency
from app.services.ssai_snapshot_repository import SnapshotReadResult
from app.services.dashboard_risk_detail_export import build_dashboard_inventory_detail_excel_bytes
from app.sims.views.dashboard_lite import (
    _build_integrated_inventory_detail_frame,
    _integrated_inventory_detail_display_frame,
)
from app.ui.chat_middleware import _build_order_calculation_analysis_profile, _build_sims_analysis_context_from_df
from tools import dashboard_display_compare_20261011 as display_compare


def check_flow() -> None:
    detail = pd.DataFrame([{
        "정렬일자": "20261011", "정렬순번": 1, "내부방향": "OUT",
        "입출고일자": "20261011", "재고위치": "첫 거래 창고",
        "재고증감": -3, "입고수량": 0, "출고수량": 3,
        "할증": 0, "공급가액": 0, "부가세": 0, "합계금액": 0,
    }])
    settings = {"stock_mode": "real", "date_basis": "io", "flow_scope": "all"}
    cases = (
        ({}, "전체 위치 합계"),
        ({"stock_cds": ["00001", "00008"]}, "선택 위치 합계"),
        ({"stock_cds": ["00001"], "stock_names": ["본사창고"]}, "본사창고"),
        ({"stock_cds": ["00001"]}, "00001"),
    )
    for params, expected in cases:
        for source in (detail, detail.iloc[:0]):
            display, meta = _prepare_display_df(source, 10, settings, params)
            assert display.iloc[0]["재고위치"] == expected
            assert display.iloc[0]["재고수량"] == 10
            assert meta["stock_qty"] == (7 if len(source) else 10)
            if len(source):
                assert display.iloc[1]["재고위치"] == "첫 거래 창고"
                assert display.iloc[1]["재고수량"] == 7
            excel = BytesIO()
            display.to_excel(excel, index=False)
            excel.seek(0)
            sheet = openpyxl.load_workbook(excel, read_only=True, data_only=True).active
            stock_column = list(display.columns).index("재고위치") + 1
            assert sheet.cell(2, stock_column).value == expected


def check_dashboard() -> None:
    rows = [{
        "product_code": code, "product_name": code, "규격": f"규격-{code}",
        "current_stock_qty": 10, "current_stock_amt": 100 + index,
        "stock_valuation_unit_price": 10 + index,
        "inventory_current_stock_present": True,
        "evaluation_expected_demand_qty": 10,
        "evaluation_expected_demand_present": True,
        "last_normal_inbound_date": "20261001", "inbound_data_status": "normal",
        "recent_inbound_vendor_code": f"V{index}",
        "recent_inbound_vendor_name": f"매입처-{code}",
        "recent_inbound_vendor_source": "actual_inbound",
    } for index, code in enumerate(("P1", "P2"))]
    rows.append({
        "product_code": "P3", "product_name": "P3", "current_stock_qty": 0,
        "inventory_current_stock_present": True,
        "evaluation_expected_demand_qty": 0,
        "evaluation_expected_demand_present": False,
    })
    result = _attach_inventory_status_and_frequency(
        rows, frequency_snapshot=SnapshotReadResult(status="missing"), monthly_business_days=20,
    )
    inventory = {
        "inventory_status_detail_rows": result["detail_rows"],
        "risk_detail_rows": [{"제품코드": "P1", "위험상태": "부족 주의", "위험사유": "fixture"}],
    }
    frame = _build_integrated_inventory_detail_frame(inventory)
    assert len(frame) == len(rows) == 3
    assert frame["제품코드"].is_unique
    by_code = frame.set_index("제품코드")
    assert by_code.loc["P2", "위험 품목 여부"] == "비위험 품목"
    assert by_code.loc["P2", "규격"] == "규격-P2"
    assert by_code.loc["P2", "현재재고금액"] == 101
    assert by_code.loc["P2", "재고평가단가"] == 11
    assert by_code.loc["P2", "최근입고 대표매입처명"] == "매입처-P2"
    assert by_code.loc["P2", "최근입고 대표매입처코드"] == "V1"
    assert by_code.loc["P2", "최근 정상 입고일"] == "20261001"
    assert by_code.loc["P2", "현재재고수량"] == 10
    assert by_code.loc["P2", "위험사유"] != "fixture"
    assert pd.isna(by_code.loc["P3", "현재재고금액"])
    assert pd.isna(by_code.loc["P3", "재고평가단가"])
    assert by_code.loc["P3", "최근입고 대표매입처명"] == ""
    screen = _integrated_inventory_detail_display_frame(frame, 3)
    assert screen.loc[screen["제품코드"].eq("P2"), "최근입고 대표매입처명"].iloc[0] == "매입처-P2"
    payload, info = build_dashboard_inventory_detail_excel_bytes(frame, [])
    sheet = openpyxl.load_workbook(BytesIO(payload), read_only=True, data_only=True).active
    columns = [cell.value for cell in sheet[1]]
    assert info["export_rows"] == 3
    assert "현재재고금액" in columns and "최근입고 대표매입처명" in columns
    assert sheet.cell(3, columns.index("최근입고 대표매입처명") + 1).value == "매입처-P2"

    bulk = _build_integrated_inventory_detail_frame({
        "inventory_status_detail_rows": [
            {"제품코드": f"B{index:05d}", "현재재고수량": 1}
            for index in range(24358)
        ],
        "risk_detail_rows": [
            {"제품코드": f"B{index:05d}", "위험상태": "부족 주의"}
            for index in range(9184)
        ],
    })
    assert len(bulk) == 24358 and bulk["제품코드"].is_unique
    assert bulk["위험 품목 여부"].eq("위험 품목").sum() == 9184
    assert bulk["위험 품목 여부"].eq("비위험 품목").sum() == 15174


def check_order_context() -> None:
    rows = []
    for index in range(24):
        unpriced = index < 22
        recommended = 5 if index < 12 else 0
        rows.append({
            "제품코드": f"P{index:03d}", "제품명": f"제품-{index}", "발주처": "거래처",
            "추천 발주수량": recommended, "실제 발주수량": recommended,
            "발주단가": None if unpriced else 10,
            "단가판정": "REVIEW" if unpriced else "확정",
            "발주금액(부가세포함)": None if unpriced else 11 * recommended,
            "3개월반품공급가액": 999999,
            "추천대비수정수량": 0,
        })
    rows[22].update({"추천 발주수량": 2, "실제 발주수량": 2, "발주금액(부가세포함)": 22})
    rows[23].update({"추천 발주수량": 3, "실제 발주수량": 4, "발주금액(부가세포함)": 44, "추천대비수정수량": 1})
    df = pd.DataFrame(rows)
    profile = _build_order_calculation_analysis_profile(df)
    assert profile["price_unconfirmed_rows"] == 22
    assert profile["price_unconfirmed_positive_recommended_rows"] == 12
    assert Decimal(profile["amount_confirmed_partial_sum"]) == 66
    assert profile["amount_confirmed_rows"] == 2
    assert profile["user_edited_rows"] == 1
    assert profile["actual_quantity_sum"] == profile["recommended_quantity_sum"] + 1
    ctx = _build_sims_analysis_context_from_df(df, result={}, action_name="발주 계산", params={}, meta={})
    assert ctx["order_calculation_profile"] == profile
    assert ctx["sales_time_profile"] == ctx["sales_group_profile"] == {}
    assert "3개월반품공급가액" not in ctx["sample_records"][0]
    assert "0원" in ctx["analysis_text"]
    no_price = _build_order_calculation_analysis_profile(df.iloc[:22])
    assert no_price["amount_confirmed_partial_sum"] is None


def check_typed_dashboard_export() -> None:
    frame = pd.DataFrame([
        {
            "제품코드": "00001", "규격": "10T", "현재재고금액": Decimal("100.00"),
            "재고평가단가": Decimal("12.50"), "최근 정상 입고일": "20261001",
            "최근입고 대표매입처코드": "00007", "최근입고 대표매입처명": "fixture vendor",
            "수요급증여부": False,
        },
        {
            "제품코드": "00002", "규격": "", "현재재고금액": Decimal("0"),
            "재고평가단가": Decimal("-1.25"), "최근 정상 입고일": "",
            "최근입고 대표매입처코드": "", "최근입고 대표매입처명": "",
            "수요급증여부": False,
        },
        {
            "제품코드": "00003", "규격": None, "현재재고금액": None,
            "재고평가단가": float("nan"), "최근 정상 입고일": None,
            "최근입고 대표매입처코드": None, "최근입고 대표매입처명": None,
            "수요급증여부": False,
        },
    ])
    result = display_compare.compare_dashboard_sample(frame)
    assert result["status"] == "PASS", result["mismatches"]
    assert result["rows"] == result["excel_rows"] == 3
    assert any(item["column"] == "최근 정상 입고일" for item in result["format_only"])
    assert all(item["product_code"].startswith("0000") for item in result["format_only"])

    normal_display = display_compare._integrated_inventory_detail_display_frame

    def corrupted_display(data, limit):
        shown = normal_display(data, limit)
        shown.loc[shown.index[0], "현재재고금액"] = 101
        return shown

    with patch.object(display_compare, "_integrated_inventory_detail_display_frame", side_effect=corrupted_display):
        bad = display_compare.compare_dashboard_sample(frame)
    assert bad["status"] == "FAIL"
    assert bad["mismatches"][0]["column"] == "현재재고금액"
    assert bad["mismatches"][0]["classification"] == "value_difference"
    assert set(bad["mismatches"][0]) >= {"product_code", "source", "screen", "excel"}


def main() -> None:
    with patch.object(socket.socket, "connect", side_effect=AssertionError("DB/network access forbidden")):
        check_flow()
        check_dashboard()
        check_order_context()
        check_typed_dashboard_export()
    print("display_context_fixes offline PASS: flow, dashboard, order context, typed export; DB calls=0")


if __name__ == "__main__":
    main()

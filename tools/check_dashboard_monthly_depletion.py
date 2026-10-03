"""Offline gate for Dashboard monthly depletion facts, detail, and Excel."""
from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path
import sys

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.services.business_calendar_service import BusinessDayMonthContext
from app.services.dashboard_lite_facts import (
    _attach_inventory_status_and_frequency,
    _attach_monthly_depletion_facts,
    _monthly_depletion_cross_distribution,
)
from app.services.dashboard_inventory_frequency_snapshot import FrequencyProjectionReadResult
from app.services.dashboard_risk_detail_export import build_dashboard_inventory_detail_excel_bytes
from app.services.ssai_snapshot_repository import SnapshotReadResult
from app.sims.views.dashboard_lite import (
    _build_integrated_inventory_detail_frame,
    _filter_integrated_inventory_detail_rows,
)


def _context() -> BusinessDayMonthContext:
    return BusinessDayMonthContext(
        evaluation_date="20260930",
        evaluation_month="202609",
        calendar_days_total=30,
        elapsed_calendar_days=30,
        calendar_progress_ratio=1.0,
        business_days_total=20,
        elapsed_business_days=20,
        business_day_progress_ratio=1.0,
        authority_status="ready",
    )


def _row(code: str, current: float, forecast: float, stock: float, pending: float) -> dict:
    return {
        "product_code": code,
        "product_name": code,
        "당월현재출고수량": current,
        "당월기준예상출고수량": forecast,
        "current_stock_qty": stock,
        "raw_pending_inbound_qty": pending,
        "inventory_current_stock_present": True,
        "evaluation_expected_demand_qty": forecast,
        "evaluation_expected_demand_present": True,
    }


def _attach_status(rows: list[dict]) -> dict:
    return _attach_inventory_status_and_frequency(
        rows,
        frequency_snapshot=SnapshotReadResult(status="missing", reason="fixture"),
        previous_frequency_projection=FrequencyProjectionReadResult(status="not_requested", reason="fixture"),
        monthly_business_days=20,
        pending_available=True,
        pending_authority="fixture",
    )


def main() -> int:
    rows = [
        _row("unexpected", 2, 0, 0, 0),
        _row("no-demand", 0, 0, 0, 0),
        _row("exceeded", 11, 10, 0, 0),
        _row("achieved", 10, 10, 0, 0),
        _row("stock", 4, 10, 6, 0),
        _row("pending", 4, 10, 2, 4),
        _row("short", 4, 10, 2, 1),
        _row("negative", 4, 10, -2, -3),
    ]
    summary = _attach_monthly_depletion_facts(
        rows,
        evaluation_month="202609",
        policy_date="20260930",
        pending_available=True,
        business_day_context=_context(),
        today=date(2026, 9, 30),
    )
    expected = {
        "unexpected": "예상외 출고",
        "no-demand": "월예상 수요 없음",
        "exceeded": "월예상 초과 달성",
        "achieved": "월예상 달성",
        "stock": "재고로 월말 충족 가능",
        "pending": "입고예정 반영 시 충족",
        "short": "월말 부족 예상",
        "negative": "월말 부족 예상",
    }
    assert summary["evaluation_status"] == "진행중"
    assert sum(summary["status_counts"].values()) == len(rows)
    assert {row["product_code"]: row["월간 소진상태"] for row in rows} == expected
    assert rows[-1]["입고예정수량"] == 0 and rows[-1]["입고예정 반영 가용재고"] == 0
    assert rows[4]["월잔여예상수요"] == 6 and rows[4]["월예상 달성률"] == 40
    assert all(row["경과영업일"] == 20 and row["전체영업일"] == 20 for row in rows)

    unavailable = [_row("unavailable", 4, 10, 2, 100)]
    unavailable_summary = _attach_monthly_depletion_facts(
        unavailable,
        evaluation_month="202609",
        policy_date="20260930",
        pending_available=False,
        business_day_context=_context(),
        today=date(2026, 9, 30),
    )
    assert unavailable[0]["월간 소진상태"] == "월말 부족 예상"
    assert unavailable[0]["입고예정수량"] is None
    assert unavailable[0]["입고예정 반영 가용재고"] is None
    assert unavailable_summary["additional_source_call_count"] == 0

    completed = [_row("complete", 4, 10, 100, 100)]
    completed_summary = _attach_monthly_depletion_facts(
        completed,
        evaluation_month="202609",
        policy_date="20260930",
        pending_available=True,
        business_day_context=_context(),
        today=date(2026, 10, 1),
    )
    assert completed_summary["evaluation_status"] == "완료월"
    assert completed[0]["월간 소진상태"] == "월예상 미달 종료"
    assert completed[0]["월간 소진상태"] not in {
        "재고로 월말 충족 가능", "입고예정 반영 시 충족", "월말 부족 예상",
    }

    before_status = [_row("A", 1, 10, 1, 0), _row("B", 1, 10, 100, 0)]
    control = [dict(row) for row in before_status]
    expected_inventory = _attach_status(control)["summary"]["status_counts"]
    _attach_monthly_depletion_facts(
        before_status,
        evaluation_month="202609",
        policy_date="20260930",
        pending_available=True,
        business_day_context=_context(),
        today=date(2026, 9, 30),
    )
    status_result = _attach_status(before_status)
    assert status_result["summary"]["status_counts"] == expected_inventory
    cross = _monthly_depletion_cross_distribution(before_status)
    assert sum(int(row["품목수"]) for row in cross) == len(before_status)

    detail_rows = status_result["detail_rows"]
    assert all(row.get("월간 소진상태") for row in detail_rows)
    assert all(row.get("당월현재출고수량") is not None for row in detail_rows)
    assert all(row.get("당월기준예상출고수량") is not None for row in detail_rows)
    inventory = {
        "inventory_status_detail_rows": detail_rows,
        "risk_detail_rows": [{
            "제품코드": "A", "위험상태": "긴급 부족",
            "당월현재출고수량": 999, "당월기준예상출고수량": 999,
            "월간 소진상태": "덮어쓰기 금지",
        }],
    }
    integrated = _build_integrated_inventory_detail_frame(inventory)
    product_a = integrated.loc[integrated["제품코드"].eq("A")].iloc[0]
    assert product_a["당월현재출고수량"] == 1
    assert product_a["당월기준예상출고수량"] == 10
    assert product_a["월간 소진상태"] == "월말 부족 예상"
    filtered = _filter_integrated_inventory_detail_rows(
        integrated,
        inventory_status="전체",
        frequency_grade="전체",
        risk_filter="전체",
        vendor_key="전체",
        search_text="",
        monthly_depletion_status="월말 부족 예상",
    )
    assert set(filtered["제품코드"]) == {"A"}

    payload, info = build_dashboard_inventory_detail_excel_bytes(integrated, [{"조건명": "fixture", "값": "offline"}])
    assert info["export_rows"] == len(integrated)
    workbook = openpyxl.load_workbook(BytesIO(payload), read_only=True, data_only=True)
    sheet = workbook["재고현황상세"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    for column in (
        "월간 소진상태", "월간 소진상태 사유", "당월현재출고수량", "당월기준예상출고수량",
        "월예상 달성률", "월잔여예상수요", "입고예정수량", "입고예정 반영 가용재고",
        "평가월", "판단기준일",
    ):
        assert column in headers, column
    status_index = headers.index("월간 소진상태")
    assert all(row[status_index].value for row in sheet.iter_rows(min_row=2))

    assert sum(summary["status_counts"].values()) == len(rows)
    print(
        "PASS dashboard monthly depletion: ongoing=8 completed=1 unavailable=1 "
        "detail_rows=2 excel_rows=2 additional_source_call_count=0"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

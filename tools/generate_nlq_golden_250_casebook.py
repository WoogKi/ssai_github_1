"""Build a deterministic Golden-250 NLQ casebook from the approved workbook.

The workbook remains the immutable authority.  This tool only exports its raw
cells and records static compatibility/execution evidence for regression
orchestration; it does not call ERP or import Streamlit runtime code.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path
import re
import sys
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.sims.nlq.action_inventory import implemented_actions
from app.ui.current_table_followups.action_dispatcher import classify_current_table_followup_intent
from tools.check_nlq_official_206_regression import ACTION_COMPATIBILITY


EXPECTED_HEADERS = (
    "순번", "선행질문", "질문", "기준 기능 (action)", "처리여부", "의도일치", "비고",
)
CURRENT_TABLE_MARKERS = ("현재표", "현재 표", "현재결과", "현재 결과")


def _cell(value: Any) -> str:
    return "" if value is None else str(value)


def _normalized(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _stable_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _sha256(value: Any) -> str:
    return hashlib.sha256(_stable_json(value).encode("utf-8")).hexdigest()


def _read_source(path: Path) -> tuple[list[dict[str, str]], str]:
    workbook_bytes = path.read_bytes()
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "NLQ 사례" not in workbook.sheetnames:
        raise ValueError(f"NLQ 사례 sheet missing: {workbook.sheetnames!r}")
    sheet = workbook["NLQ 사례"]
    values = list(sheet.iter_rows(values_only=True))
    headers = tuple(_normalized(_cell(value)) for value in values[0])
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"unexpected headers: {headers!r}")
    rows: list[dict[str, str]] = []
    for worksheet_row, raw in enumerate(values[1:], start=2):
        if not any(value not in (None, "") for value in raw):
            continue
        record = {headers[index]: _cell(raw[index]) for index in range(len(headers))}
        record["worksheet_row"] = str(worksheet_row)
        rows.append(record)
    return rows, hashlib.sha256(workbook_bytes).hexdigest()


def _action_catalog() -> tuple[set[str], dict[str, str]]:
    canonical: set[str] = set()
    aliases: dict[str, str] = {}
    for spec in implemented_actions():
        canonical.add(spec.canonical_action)
        for label in (*spec.label_aliases, *spec.nlq_aliases, spec.panel_action):
            normalized = _normalized(str(label))
            if normalized and normalized != spec.canonical_action:
                aliases[normalized] = spec.canonical_action
    return canonical, aliases


def _compatibility(expected: str, canonical: set[str], aliases: dict[str, str]) -> tuple[str, str, str]:
    expected = _normalized(expected)
    if expected in canonical:
        return expected, "registry.canonical_action", "EXACT"
    if expected in aliases:
        return aliases[expected], "registry.label_or_nlq_alias", "DECLARED_ALIAS"
    compatible = ACTION_COMPATIBILITY.get(expected)
    if compatible:
        return sorted(compatible)[0], "tools.check_nlq_official_206_regression.ACTION_COMPATIBILITY", "SCREEN_TO_RUNTIME_COMPATIBLE"
    return "", "", "UNDECLARED"


def _execution_mode(parent: str, question: str, expected_action: str) -> tuple[str, str]:
    if parent:
        intent = classify_current_table_followup_intent(question)
        if intent == "llm_analysis":
            return "CURRENT_TABLE_LLM", "action_dispatcher.classify_current_table_followup_intent=llm_analysis"
        return "CURRENT_TABLE_DETERMINISTIC", f"action_dispatcher.classify_current_table_followup_intent={intent}"
    if expected_action in {"발주 계산", "SIMS 일일점검"}:
        return "DIRECT_STRUCTURED", "registered action contract"
    if expected_action in {"제품정보 조회"}:
        return "DATA_DEPENDENT", "Snapshot-backed product information"
    if any(marker in question for marker in CURRENT_TABLE_MARKERS):
        return "SOURCE_REVIEW", "current-table marker without authoritative parent"
    return "DIRECT_TABLE", "no predecessor question"


def _legacy_map(old_rows: list[dict[str, str]], cases: list[dict[str, str]]) -> list[dict[str, str]]:
    old_by_pair = {(_normalized(row["선행질문"]), _normalized(row["질문"])): row for row in old_rows}
    new_by_pair = {(_normalized(row["parent_question_raw"]), _normalized(row["question_raw"])): row for row in cases}
    out: list[dict[str, str]] = []
    matched_old: set[str] = set()
    for item in cases:
        pair = (_normalized(item["parent_question_raw"]), _normalized(item["question_raw"]))
        old = old_by_pair.get(pair)
        if old:
            old_id = f"NLQ-{int(old['순번']):04d}"
            matched_old.add(old_id)
            same_action = _normalized(old["기준 기능 (action)"]) == _normalized(item["expected_action_raw"])
            out.append({
                "legacy_case_id": old_id, "golden_case_id": item["golden_case_id"],
                "old_parent": old["선행질문"], "new_parent": item["parent_question_raw"],
                "old_question": old["질문"], "new_question": item["question_raw"],
                "old_expected_action": old["기준 기능 (action)"], "new_expected_action": item["expected_action_raw"],
                "change_type": "UNCHANGED" if same_action else "ACTION_CHANGED",
                "reason": "normalized parent+question exact match",
            })
            continue
        candidates = [row for row in old_rows if _normalized(row["질문"]) == _normalized(item["question_raw"])]
        candidates = [row for row in candidates if _normalized(row["기준 기능 (action)"]) == _normalized(item["expected_action_raw"])]
        if len(candidates) == 1:
            old = candidates[0]
            old_id = f"NLQ-{int(old['순번']):04d}"
            matched_old.add(old_id)
            out.append({
                "legacy_case_id": old_id, "golden_case_id": item["golden_case_id"],
                "old_parent": old["선행질문"], "new_parent": item["parent_question_raw"],
                "old_question": old["질문"], "new_question": item["question_raw"],
                "old_expected_action": old["기준 기능 (action)"], "new_expected_action": item["expected_action_raw"],
                "change_type": "PARENT_CHANGED", "reason": "unique same question+action",
            })
        else:
            out.append({
                "legacy_case_id": "", "golden_case_id": item["golden_case_id"],
                "old_parent": "", "new_parent": item["parent_question_raw"],
                "old_question": "", "new_question": item["question_raw"],
                "old_expected_action": "", "new_expected_action": item["expected_action_raw"],
                "change_type": "NEW", "reason": "no proven legacy semantic match",
            })
    for old in old_rows:
        old_id = f"NLQ-{int(old['순번']):04d}"
        if old_id not in matched_old:
            out.append({
                "legacy_case_id": old_id, "golden_case_id": "", "old_parent": old["선행질문"],
                "new_parent": "", "old_question": old["질문"], "new_question": "",
                "old_expected_action": old["기준 기능 (action)"], "new_expected_action": "",
                "change_type": "REMOVED", "reason": "no proven Golden semantic match",
            })
    return out


def _write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(rows[0]) if rows else []
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def generate(
    source: Path,
    legacy: Path,
    output_dir: Path,
    *,
    id_prefix: str = "G250",
    file_prefix: str = "nlq_golden_250",
    expected_count: int = 250,
) -> dict[str, Any]:
    source_rows, source_checksum = _read_source(source)
    seqs = [int(row["순번"]) for row in source_rows]
    if expected_count < 1 or len(source_rows) != expected_count or seqs != list(range(1, expected_count + 1)):
        raise ValueError(f"Golden source must contain exactly sequential rows 1..{expected_count}")
    pairs = [(_normalized(row["선행질문"]), _normalized(row["질문"])) for row in source_rows]
    if any(not _normalized(row["질문"]) or not _normalized(row["기준 기능 (action)"]) for row in source_rows):
        raise ValueError("Golden source has blank question/action")
    if len(set(pairs)) != len(pairs):
        raise ValueError("Golden source has duplicate normalized parent+question")
    canonical, aliases = _action_catalog()
    cases: list[dict[str, str]] = []
    for row in source_rows:
        expected = row["기준 기능 (action)"]
        resolved, source_name, status = _compatibility(expected, canonical, aliases)
        mode, evidence = _execution_mode(row["선행질문"], row["질문"], expected)
        cases.append({
            "golden_case_id": f"{id_prefix}-{int(row['순번']):04d}", "source_seq": row["순번"],
            "parent_question_raw": row["선행질문"], "question_raw": row["질문"],
            "expected_action_raw": expected, "source_process_status": row["처리여부"],
            "historical_intent_result": row["의도일치"], "note_raw": row["비고"],
            "worksheet_row": row["worksheet_row"], "resolved_expected_action": resolved,
            "compatibility_source": source_name, "compatibility_status": status,
            "execution_mode": mode, "execution_mode_evidence": evidence,
        })
    old_rows, _ = _read_source(legacy)
    mapping = _legacy_map(old_rows, cases)
    legacy_by_golden = {item["golden_case_id"]: item["legacy_case_id"] for item in mapping if item["golden_case_id"] and item["legacy_case_id"]}
    for case in cases:
        case["legacy_case_id"] = legacy_by_golden.get(case["golden_case_id"], "")
    compatibility = [{key: item[key] for key in (
        "golden_case_id", "source_seq", "question_raw", "expected_action_raw",
        "resolved_expected_action", "compatibility_source", "compatibility_status",
    )} for item in cases]
    execution = [{key: item[key] for key in (
        "golden_case_id", "source_seq", "parent_question_raw", "question_raw",
        "expected_action_raw", "execution_mode", "execution_mode_evidence",
    )} for item in cases]
    _write_csv(output_dir / f"{file_prefix}_casebook_20261007.csv", cases)
    _write_csv(output_dir / f"{file_prefix}_legacy_227_map_20261007.csv", mapping)
    _write_csv(output_dir / f"{file_prefix}_action_compatibility_20261007.csv", compatibility)
    _write_csv(output_dir / f"{file_prefix}_execution_modes_20261007.csv", execution)
    payload = {"source": {"path": str(source), "checksum_sha256": source_checksum, "sheet": "NLQ 사례"}, "cases": cases}
    (output_dir / f"{file_prefix}_casebook_20261007.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return {
        "source_checksum_sha256": source_checksum, "casebook_checksum_sha256": _sha256(cases),
        "source_rows": len(cases), "compatibility": dict(Counter(item["compatibility_status"] for item in cases)),
        "execution_modes": dict(Counter(item["execution_mode"] for item in cases)),
        "legacy_changes": dict(Counter(item["change_type"] for item in mapping)),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=ROOT / "docs" / "NLQ 정리 20261007.xlsx")
    parser.add_argument("--legacy", type=Path, default=ROOT / "docs" / "NLQ 정리 20260807.xlsx")
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--id-prefix", default="G250")
    parser.add_argument("--file-prefix", default="nlq_golden_250")
    parser.add_argument("--expected-count", type=int, default=250)
    args = parser.parse_args()
    print(json.dumps(generate(
        args.source, args.legacy, args.output_dir,
        id_prefix=args.id_prefix, file_prefix=args.file_prefix, expected_count=args.expected_count,
    ), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

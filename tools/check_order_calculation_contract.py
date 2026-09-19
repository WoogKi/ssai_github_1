"""Offline contract and production-adapter/editor fixtures; never connects to DB."""
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path
import sys
from unittest.mock import patch
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.services.order_calculation_contract import (
    OrderConditions, horizon_dates, demand_for_dates, demand_trend_adjustment,
    calculate_quantities, recommend_quantity, amounts,
)
from app.services.order_calculation_service import assemble_result, get_order_calculation_result
from app.ui.order_calculation_editor import apply_actual_edits


def test_horizon():
    start = date(2026, 9, 25)
    dates = [start + timedelta(days=i) for i in range(-25, 70) if (start + timedelta(days=i)).weekday() < 5]
    before = horizon_dates(OrderConditions(start), dates)
    assert len(before) == 15 and min(before) == date(2026, 9, 28)
    assert len(horizon_dates(OrderConditions(date(2026, 9, 28)), dates)) == 15
    assert len(horizon_dates(OrderConditions(start, 0, 1, 0), dates)) == 1
    assert len(horizon_dates(OrderConditions(start, 0, 15, 0), dates)) == 15


def test_horizon_settlement_boundary():
    # The authority dates include holiday exclusions; never infer by subtraction.
    excluded = {date(2026, 9, 24), date(2026, 9, 25)}
    dates = [date(2026, 9, 1) + timedelta(days=i) for i in range(90)
             if (date(2026, 9, 1) + timedelta(days=i)).weekday() < 5
             and date(2026, 9, 1) + timedelta(days=i) not in excluded]
    for day, expected in ((1, 15), (10, 12), (14, 10), (25, 15), (28, 15), (30, 15)):
        reference = date(2026, 9, day)
        selected = horizon_dates(OrderConditions(reference), dates)
        assert len(selected) == expected and all(d > reference for d in selected)
        if day < 25:
            remaining = tuple(d for d in dates if d > reference and d.month == 9)
            assert selected == remaining[:15]
    assert horizon_dates(OrderConditions(date(2026, 9, 12)), dates)[0] == date(2026, 9, 14)
    assert len(horizon_dates(OrderConditions(date(2026, 9, 25), closing_day=0), dates)) == 15
    assert len(horizon_dates(OrderConditions(date(2026, 9, 25), closing_day=30), dates)) == 3


def test_monthly_allocation():
    reference = date(2026, 9, 25)
    dates = [date(2026, 9, 28), date(2026, 9, 29), date(2026, 9, 30), date(2026, 10, 1), date(2026, 10, 2)]
    value, missing = demand_for_dates(dates[:2], reference_date=reference, business_dates=dates, monthly_forecast={"202609": D(30)})
    assert value == 20 and not missing
    value, missing = demand_for_dates(dates, reference_date=reference, business_dates=dates, monthly_forecast={"202609": D(30)})
    assert value is None and missing == ("202610",)
    value, missing = demand_for_dates(dates, reference_date=reference, business_dates=dates, monthly_forecast={"202609": D(30), "202610": D(40)})
    assert value == 70 and not missing


def test_quantities():
    result = calculate_quantities(stock=D(-2), pending=D(10), safety_demand=D(3), horizon_demand=D(50))
    assert result["계산 발주수량"] == 42 and result["추천 발주수량"] == 42
    assert recommend_quantity(D(43), D(10), increasing=True) == 50
    assert recommend_quantity(D(43), D(10), increasing=False) == 40
    assert recommend_quantity(D(3), D(10), increasing=False) == 10
    assert recommend_quantity(D('6.56'), D(10), increasing=False) == 10
    assert recommend_quantity(D('12.365'), D(100), increasing=False) == 100
    assert recommend_quantity(D(26), D(10), increasing=False) == 20
    assert recommend_quantity(D(26), D(10), increasing=True) == 30
    assert calculate_quantities(stock=D(100), pending=D(0), safety_demand=D(3), horizon_demand=D(150), unit=D(100))['추천 발주수량'] == 0
    assert recommend_quantity(D(3), D(10), increasing=True) == 10
    for raw, expected_up, expected_down in (("12.4", 13, 12), ("0.7", 1, 0)):
        assert recommend_quantity(D(raw), None, increasing=True) == expected_up
        assert recommend_quantity(D(raw), None, increasing=False) == expected_down
    try:
        amounts(D("1.2"), D(100))
        raise AssertionError("fractional actual quantity accepted")
    except ValueError:
        pass
    assert recommend_quantity(D(43), None, increasing=False) == 43
    assert calculate_quantities(stock=D(100), pending=D(0), safety_demand=D(3), horizon_demand=D(50))["추천 발주수량"] == 0


def test_demand_trend_adjustment():
    cases = (
        (0, 0), (5, 0), (10, 0), (20, 10), (40, 20), (60, 30), (100, 30),
        (-5, 0), (-10, 0), (-20, -10), (-40, -20), (-60, -30), (-100, -30),
    )
    for rate_pct, expected_pct in cases:
        recent = D(100) * (D(1) + D(rate_pct) / D(100))
        rate, adjustment, _reason = demand_trend_adjustment(
            recent_3m_avg=recent, previous_3m_avg=D(100), completed_months=6, frequency_grade="A",
        )
        assert rate == D(rate_pct) / D(100)
        assert adjustment == D(expected_pct) / D(100)
    for grade in ("F", "X"):
        rate, adjustment, reason = demand_trend_adjustment(
            recent_3m_avg=D(0), previous_3m_avg=D(100), completed_months=6, frequency_grade=grade,
        )
        assert rate is None and adjustment == 0 and grade in reason
    assert demand_trend_adjustment(
        recent_3m_avg=D(2), previous_3m_avg=D(1), completed_months=5, frequency_grade="A",
    )[1] == 0

    boundary_cases = (
        ("90.0001", "0"), ("90.0000", "0"), ("89.9999", "-0.0500005"),
        ("109.9999", "0"), ("110.0000", "0"), ("110.0001", "0.0500005"),
    )
    for recent_text, expected_adjustment in boundary_cases:
        _rate, adjustment, _reason = demand_trend_adjustment(
            recent_3m_avg=D(recent_text), previous_3m_avg=D(100),
            completed_months=6, frequency_grade="A",
        )
        assert adjustment == D(expected_adjustment), (recent_text, adjustment)

    # Product 27252 equivalent: the recurring three-month average arrives as
    # a finite decimal just below 13 1/3, but the business change is -10%.
    for previous_text in ("13.33333333333333", "13.333333333333332"):
        _rate, adjustment, reason = demand_trend_adjustment(
            recent_3m_avg=D(12), previous_3m_avg=D(previous_text),
            completed_months=6, frequency_grade="A",
        )
        assert adjustment == 0 and reason == "deadband", previous_text


def test_full_month_forecast_and_blank_pending():
    from app.services.order_calculation_service import build_pending_params
    params, sources = fixture()
    params["order_date"] = "2026-09-14"
    sources["business_dates"] = [date(2026, 9, d) for d in range(1, 31)
                                 if date(2026, 9, d).weekday() < 5]
    sources["demand"]["당월 예상출고수량"] = D("85.46")
    sources["demand"]["당월 현재출고수량"] = D(40)
    sources["demand"]["당월 잔여예상출고수량"] = D("45.46")
    sources["demand"]["현재재고수량"] = D(12)
    sources["pending"] = pd.DataFrame({"제품코드": ["00001"], "입고예정수량": [D(10)]})
    sources["elapsed_days"] = 10
    sources["normal_outbound_verified"] = True
    row = assemble_result(params, sources).iloc[0]
    assert row["해당 월 전체 영업일수"] == 22
    assert row["수요근거"] == "예측기반"
    assert row["기준 1영업일 예상수량"] == D("85.46") / 22
    assert row["계산 발주수량"] == row["horizon 예정수량"] - 12 - 10
    assert row["적용 필요 영업일수"] == row["적용 horizon 영업일수"]
    assert row["적용 필요예정수량"] == row["horizon 필요예정수량"]
    sources["demand"]["당월 현재출고수량"] = D(80)
    changed = assemble_result(params, sources).iloc[0]
    assert changed["horizon 예정수량"] == row["horizon 예정수량"]
    sources["demand"]["수요예상기준"] = "비교자료부족"
    pace = assemble_result(params, sources).iloc[0]
    assert pace["수요근거"] == "실적기반" and pace["기준 1영업일 예상수량"] == 8
    assert pace["horizon 필요예정수량"] == 96 and pace["안전재고 기준수량"] == 24
    assert pace["계산 발주수량"] == 74 and pace["추천 발주수량"] == 74
    sources["demand"]["당월 현재출고수량"] = D(0)
    missing = assemble_result(params, sources).iloc[0]
    assert missing["수요근거"] == "수요근거 없음" and missing["추천 발주수량"] is None
    q = build_pending_params({"policy_date": "20260914", "stock_cd_list": ["00001"],
                              "order_vendor_cd": "different", "date_to": "20260914"})
    assert q["include_blank_stock_cd"] and "date_to" not in q and "order_vendor_cd" not in q
    from app.services import rddbc170_rddbc180_order_service as pending
    with patch.object(pending, "execute_bound_select", return_value=pd.DataFrame()) as query:
        pending.get_expected_inbound_product_totals({**q, "_business_dates": ("20260914", "20260911", "20260910", "20260909")})
    sql = query.call_args.args[0]
    assert "OR NULLIF(RTRIM(D.Rd18_Stock_Cd), '') IS NULL" in sql
    assert " - D.Rd18_In_Quantity" in sql
    calendar = sources["business_dates"]
    for plan in (D(1), D(10), D(43)):
        total, missing = demand_for_dates(calendar, reference_date=date(2026, 8, 31),
                                         business_dates=calendar, monthly_forecast={"202609": plan})
        assert total == plan and not missing
        assert recommend_quantity(total, None, increasing=True) == plan
        assert recommend_quantity(total, None, increasing=False) == plan


def fixture():
    base = pd.DataFrame({"제품코드": ["00001"], "제품명": ["fixture"], "출고빈도등급": ["F"],
                         "품목손익등급": ["B"], "품목기여등급": ["A"],
                         "추정단위손익": [D(10)], "추정손익률": [D("0.1")], "추정기여금액": [D(300)],
                         "제품등록실단가": [D(100)]})
    demand = pd.DataFrame({"제품코드": ["00001"], "현재재고수량": [D(-2)],
                           "당월 예상출고수량": [D(30)], "수요예상기준": ["최근3개월평균수요수량"]})
    supplier = pd.DataFrame({"product_code": ["00001"], "recent_inbound_vendor_code": ["00100"],
                             "recent_inbound_vendor_name": ["fixture-vendor"],
                             "recent_inbound_vendor_count_90": [2],
                             "recent_inbound_vendor_staff_code": ["U001"],
                             "recent_inbound_vendor_staff_name": ["홍길동"],
                             "manufacturer_vendor_code": ["M001"],
                             "manufacturer_vendor_name": ["fixture-maker"],
                             "manufacturer_staff_code": ["P001"],
                             "manufacturer_staff_name": ["이기재"],
                             "master_order_staff_code": ["U999"],
                             "master_order_staff_name": ["다른담당자"]})
    params = {"company_id": 7, "order_date": "2026-09-25", "safety_days": 3, "target_days": 15, "closing_day": 25, "price_basis": "real"}
    sources = {"snapshot": {"meta": {"snapshot_status": "ready"}}, "base": base, "demand": demand,
        "suppliers": supplier, "pending": pd.DataFrame(), "prices": {"df": pd.DataFrame()},
        "calendar_status": "ready", "business_dates": [date(2026, 9, d) for d in (28, 29, 30)]}
    return params, sources


def test_trend_applies_before_stock_pending_and_unit():
    params, sources = fixture()
    params.update(order_date="2026-09-01", safety_days=15, target_days=15)
    sources["business_dates"] = [date(2026, 9, day) for day in range(2, 31)
                                 if date(2026, 9, day).weekday() < 5][:15]
    sources["demand"]["현재재고수량"] = D(60)
    sources["demand"]["당월 예상출고수량"] = D(999)
    sources["demand"]["예상기준월수량"] = D(100)
    sources["demand"]["최근3개월평균수요수량"] = D(140)
    sources["demand"]["직전3개월평균수요수량"] = D(100)
    sources["demand"]["완료월수"] = 6
    sources["base"]["출고빈도등급"] = "A"
    sources["pending"] = pd.DataFrame({"제품코드": ["00001"], "입고예정수량": [D(20)]})
    row = assemble_result(params, sources).iloc[0]
    assert row["base_demand_qty"] == 100
    assert row["trend_rate"] == D("0.4") and row["trend_adjustment"] == D("0.2")
    assert row["adjusted_demand_qty"] == 120
    assert row["계산 발주수량"] == 40 and row["추천 발주수량"] == 40
    assert row["raw_order_qty"] == 40 and row["final_recommended_qty"] == 40

    for grade in ("F", "X"):
        grade_sources = {**sources, "base": sources["base"].copy(), "demand": sources["demand"].copy()}
        grade_sources["base"]["출고빈도등급"] = grade
        grade_row = assemble_result(params, grade_sources).iloc[0]
        assert grade_row["trend_adjustment"] == 0 and grade_row["adjusted_demand_qty"] == 100

    pace_sources = {**sources, "base": sources["base"].copy(), "demand": sources["demand"].copy()}
    pace_sources["base"]["출고빈도등급"] = "X"
    pace_sources["demand"]["예상기준월수량"] = D(0)
    pace_sources["demand"]["당월 예상출고수량"] = D(0)
    pace_sources["demand"]["당월 현재출고수량"] = D(10)
    pace_sources["elapsed_days"] = 5
    pace_sources["normal_outbound_verified"] = True
    pace_row = assemble_result(params, pace_sources).iloc[0]
    assert pace_row["수요근거"] == "실적기반"
    assert pace_row["fallback_reason"] == "current_month_actual_pace"


def test_assembly_edit_export():
    params, sources = fixture()
    frame = assemble_result(params, sources)
    assert frame.iloc[0]["출고빈도등급"] == "F" and frame.iloc[0]["제품코드"] == "00001"
    assert frame.iloc[0]["추천 발주수량"] == 32
    updated = apply_actual_edits(frame, frame, {0: {"실제 발주수량": "50"}})
    assert updated.iloc[0]["계산 발주수량"] == 32 and updated.iloc[0]["추천 발주수량"] == 32
    assert updated.iloc[0]["실제 발주수량"] == 50 and updated.iloc[0]["발주금액(부가세포함)"] == 5500
    assert frame.iloc[0]["실제 발주수량"] == 32
    assert "50" in updated.to_csv(index=False)
    try:
        apply_actual_edits(frame, frame, {0: {"실제 발주수량": "-1"}})
    except ValueError:
        pass
    else:
        raise AssertionError("negative actual quantity accepted")


def test_company_isolation():
    params, sources = fixture()
    with patch("app.services.order_calculation_service.get_current_company_id", return_value=4):
        try:
            get_order_calculation_result(params, source_loader=lambda p: (_ for _ in ()).throw(AssertionError("source called")))
        except ValueError:
            pass
        else:
            raise AssertionError("company mismatch accepted")
    with patch("app.services.order_calculation_service.get_current_company_id", return_value=7):
        result = get_order_calculation_result(params, source_loader=lambda p: sources)
    assert result["meta"]["source_call_count"] == 0 and len(result["df"]) == 1


def test_purchase_vendor_count_staff_and_sensitive_projection():
    params, sources = fixture()
    row = assemble_result(params, sources).iloc[0]
    assert row["매입거래처수"] == 2
    assert row["발주담당자코드"] == "U001" and row["발주담당자"] == "홍길동"
    assert row["제약담당자코드"] == "P001" and row["제약담당자"] == "이기재"
    assert row["발주담당자"] != sources["suppliers"].iloc[0]["master_order_staff_name"]
    assert len(assemble_result({**params, "order_staff_nm": "홍길"}, sources)) == 1
    assert assemble_result({**params, "order_staff_nm": "없는담당자"}, sources).empty
    assert len(assemble_result({**params, "pharma_staff_nm": "이기"}, sources)) == 1
    assert assemble_result({**params, "pharma_staff_nm": "없는담당자"}, sources).empty
    sensitive = {"추정단위손익", "추정손익률", "추정기여금액"}
    with patch("app.services.order_calculation_service.get_current_company_id", return_value=7), \
         patch("app.services.snapshot_product_information_service.product_information_cost_visible", return_value=False):
        member = get_order_calculation_result(params, source_loader=lambda q: sources)
    assert not sensitive.intersection(member["df"].columns)
    assert not sensitive.intersection(member["df_display"].columns)
    assert not sensitive.intersection(member["records"][0])
    with patch("app.services.order_calculation_service.get_current_company_id", return_value=7), \
         patch("app.services.snapshot_product_information_service.product_information_cost_visible", return_value=True):
        management = get_order_calculation_result(params, source_loader=lambda q: sources)
    assert sensitive.issubset(management["df"].columns)


def test_routes_menu():
    from app.services.erp_table_nlq import resolve_registered_erp_table_nlq
    from app.services.io_nlq import resolve_io_nlq
    from app.sims.meta.erp_table_feature_registry import menu_targets
    for resolver in (resolve_registered_erp_table_nlq, resolve_io_nlq):
        for text in ("발주 보여줘", "발주내역", "발주현황", "제약사 중외제약 발주 보여줘"):
            assert resolver(text)["action"] == "발주조회"
        for text in ("발주 계산", "발주 추천", "발주 수량 계산"):
            assert resolver(text)["action"] == "발주 계산"
    options = menu_targets(business_group="재고관리")
    assert options.index("발주 계산") == options.index("발주") + 1


def test_panel_submission():
    from streamlit.testing.v1 import AppTest
    script = "from app.sims.views.order_calculation_view import view_order_calculation\nview_order_calculation()"
    calls = []
    def query(params, **kwargs):
        calls.append(params)
        return {"action": "발주 계산", "final": True}
    with patch("app.sims.views.order_calculation_view.render_product_master_filters", return_value={}), \
         patch("app.sims.views.order_calculation_view.get_order_calculation_result", side_effect=query):
        app = AppTest.from_string(script).run()
        assert not app.exception and not calls
        app.button[0].click().run()
        assert not app.exception and len(calls) == 1
        assert calls[0]["safety_days"] == 3 and calls[0]["target_days"] == 15
        assert calls[0]["query_mode"] == "발주해당자료만" and calls[0]["only_needed"] and "price_basis" not in calls[0]
        assert app.session_state["__sims_panel_active"] is True
        assert app.selectbox[0].options == ["전체", "발주해당자료만", "확인 필요"]
        app.run()
        assert not app.exception and len(calls) == 1
        assert app.button[0].label == "발주 계산"
        app.button[0].click().run()
        assert len(calls) == 2


def test_price_conflict_and_field_status():
    params, sources = fixture()
    sources["prices"] = {"df": pd.DataFrame({"제품코드": ["00001", "00001"],
        "매입처코드": ["00100", "00100"], "실입고단가": [D(100), D(101)]})}
    frame = assemble_result(params, sources)
    assert frame.iloc[0]["추천 발주수량"] == 32
    assert frame.iloc[0]["단가출처"] == "복수 최종매입가 사용자확인"
    assert frame.iloc[0]["발주금액(부가세포함)"] is None
    sources["prices"]["df"]["실입고단가"] = [D(100), D(100)]
    assert assemble_result(params, sources).iloc[0]["발주단가"] == 100
    params["cost_apply_cd"] = "00007"
    sources["contracts"] = {"df": pd.DataFrame({"제품코드": ["00001"], "실입고단가": [D(200)]})}
    assert assemble_result(params, sources).iloc[0]["발주단가"] == 200
    sources["contracts"]["df"]["실입고단가"] = [D(0)]
    frame = assemble_result(params, sources)
    assert frame.iloc[0]["발주단가"] is None and frame.iloc[0]["추천 발주수량"] == 32


def test_scoped_timeout_and_no_retry():
    from types import SimpleNamespace
    import app.db.mssql_client as db
    driver = SimpleNamespace(timeout=7)
    closed = []
    connection = SimpleNamespace(connection=SimpleNamespace(driver_connection=driver), close=lambda: closed.append(True))
    engine = SimpleNamespace(connect=lambda: connection)
    def read(*args, **kwargs):
        assert driver.timeout == 120
        return pd.DataFrame({"value": [1]})
    with patch.object(db, "_get_engine", return_value=engine), patch.object(db.pd, "read_sql", side_effect=read) as sql:
        with db.read_only_request(timeout_seconds=120) as measurement:
            db.read_df("-- fixture\nSELECT 1")
        assert driver.timeout == 7 and sql.call_count == 1
        assert len(measurement["queries"]) == 1
    with patch.object(db, "_get_engine", return_value=engine), patch.object(db.pd, "read_sql", side_effect=RuntimeError("fixture")) as sql:
        with db.read_only_request(timeout_seconds=120) as measurement:
            try:
                db.read_df("SELECT 1")
            except RuntimeError:
                pass
        assert sql.call_count == 1 and driver.timeout == 7
        assert measurement["queries"][0]["status"] == "error"


def test_production_nlq_dispatch():
    import logging
    from app.sims.nlq.nlq_router import _try_handle_io_nlq
    captured = []
    payload = {"action": "발주 계산", "title": "발주 계산", "table": "order_calculation",
               "final": True, "data": "fixture", "params": {}, "meta": {"result_status": "success"}}
    with patch("app.services.order_calculation_service.get_order_calculation_result", return_value=payload) as service, \
         patch("app.ui.chat_middleware.push_sims_result_to_chat", side_effect=lambda p, a: captured.append((p, a)) or p.get("meta", {})):
        handled = _try_handle_io_nlq("발주 계산", room={}, session_state={},
            make_ts=lambda: "fixture", next_seq=lambda: 1, logger=logging.getLogger("fixture"))
    assert handled and service.call_count == 1 and len(captured) == 1
    assert captured[0][1] == "발주 계산"


def test_editor_callback_ownership_and_cache():
    from contextlib import nullcontext
    from types import SimpleNamespace
    import app.ui.order_calculation_editor as ui
    params, sources = fixture()
    full = assemble_result(params, sources)
    full.loc[full.index[0], "계산 발주수량"] = D("-12")
    state = {"sims_export_tables": {"fixture-key": full}, "sims_tables": {"fixture-key": full},
             "__sims_current_table_source_key": "fixture-key", "__sims_download_bytes::uid": b"old"}
    meta = {"table_key": "fixture-key"}
    item = {"meta": meta, "data": full}
    captured = {}
    def editor(frame, **kwargs):
        captured.update(kwargs)
        captured["frame"] = frame
    fake = SimpleNamespace(session_state=state, selectbox=lambda *a, **kw: 1,
        expander=lambda *a, **kw: nullcontext(), data_editor=editor, error=lambda *a: None,
        column_config=SimpleNamespace(TextColumn=lambda *a, **kw: None, NumberColumn=lambda *a, **kw: None))
    with patch.object(ui, "st", fake), patch("app.ui.chat_middleware._build_sims_context_from_result") as context, patch("app.db.mssql_client.get_current_company_id", return_value=7) as company:
        ui.render_actual_quantity_editor(item, meta, download_uid="uid")
        assert "실제 발주수량" not in captured["disabled"]
        assert "계산 발주수량" in captured["disabled"] and "추천 발주수량" in captured["disabled"]
        assert str(captured["frame"]["실제 발주수량"].dtype) == "Int64"
        assert pd.isna(captured["frame"].iloc[0]["계산 발주수량"])
        assert pd.api.types.is_numeric_dtype(captured['frame']['추천 발주수량'])
        state[captured["key"]] = {"edited_rows": {0: {"실제 발주수량": "55"}}}
        company.return_value = 4
        captured["on_change"]()
        assert state["sims_export_tables"]["fixture-key"].iloc[0]["실제 발주수량"] == 32
        company.return_value = 7
        captured["on_change"]()
        assert state["sims_export_tables"]["fixture-key"].iloc[0]["실제 발주수량"] == 55
        assert state["__sims_export_tables_by_key"]["fixture-key"].iloc[0]["발주금액(부가세포함)"] == 6050
        assert item["df"].iloc[0]["추천 발주수량"] == 32
        assert item["df"].iloc[0]["계산 발주수량"] == D("-12")
        assert state["__sims_current_table_source_key"] == "fixture-key"
        assert "__sims_download_bytes::uid" not in state
        assert meta["quantity_edit_revision"] == 1
        assert context.call_args.args[3].iloc[0]["실제 발주수량"] == 55
        first_key = captured["key"]
        ui.render_actual_quantity_editor(item, meta, download_uid="uid")
        assert captured["key"] == first_key
        state[first_key] = {"edited_rows": {0: {"실제 발주수량": 60}}}
        captured["on_change"]()
        assert item["df"].iloc[0]["실제 발주수량"] == 60
        assert item["df"].iloc[0]["발주금액(부가세포함)"] == 6600


def test_excel_numeric_round_trip():
    from io import BytesIO
    import openpyxl
    from app.ui.chat_middleware import _sanitize_dataframe_for_excel
    params, sources = fixture()
    original = assemble_result(params, sources)
    original.loc[original.index[0], "제품코드"] = "00123"
    original.loc[original.index[0], "계산 발주수량"] = D("-12.345")
    original.loc[original.index[0], "추정손익률"] = D("-0.123456")
    numeric = _sanitize_dataframe_for_excel(original)
    output = BytesIO()
    numeric.to_excel(output, index=False)
    sheet = openpyxl.load_workbook(BytesIO(output.getvalue())).active
    columns = {c.value: c.column for c in sheet[1]}
    for column in ("계산 발주수량", "추천 발주수량", "실제 발주수량", "재고수량", "발주단가", "발주금액(부가세포함)", "추정손익률"):
        cell = sheet.cell(2, columns[column])
        assert cell.data_type == "n", (column, cell.data_type)
    assert sheet.cell(2, columns["제품코드"]).value == "00123"
    assert sheet.cell(2, columns["제품코드"]).data_type == "s"
    assert original.iloc[0]["계산 발주수량"] == D("-12.345")


def test_production_editor_render_boundary():
    import ast
    from pathlib import Path
    tree = ast.parse(Path("app/ui/chat_middleware.py").read_text(encoding="utf-8-sig"))
    branches = [n for n in ast.walk(tree) if isinstance(n, ast.If) and isinstance(n.test, ast.Name) and n.test.id == "order_editor_rendered"]
    assert len(branches) == 1
    branch = branches[0]
    assert len(branch.body) == 1 and isinstance(branch.body[0], ast.Pass)
    assert isinstance(branch.orelse[0], ast.If) and branch.orelse[0].test.id == "is_io_table"
    code = compile(ast.Module(body=[branch], type_ignores=[]), "render-boundary", "exec")
    # Evaluate the real production branch: an editor-owned result cannot reach either general renderer.
    exec(code, {"order_editor_rendered": True})


def test_streamlit_editor_widget():
    from streamlit.testing.v1 import AppTest
    script = '''
import streamlit as st
from tools.check_order_calculation_contract import fixture, assemble_result
from app.db.mssql_client import set_current_company_id
from app.ui.order_calculation_editor import render_actual_quantity_editor
set_current_company_id(7)
p, s = fixture()
f = assemble_result(p, s)
st.session_state.setdefault('sims_export_tables', {'widget': f})
st.session_state.setdefault('sims_tables', {'widget': f})
st.session_state['__sims_current_table_source_key'] = 'widget'
render_actual_quantity_editor({'meta': {'table_key': 'widget'}}, {'table_key': 'widget'}, download_uid='widget')
'''
    at = AppTest.from_string(script).run(timeout=10)
    assert not at.exception
    assert len(at.dataframe) == 1
    proto = at.dataframe[0].proto
    assert proto.editing_mode != 0
    import json
    config = json.loads(proto.columns)
    for column in ('계산 발주수량', '추천 발주수량', '재고수량', '입고예정수량', '발주단가', '발주금액(부가세포함)', '월 기준 예상수량', '안전재고 기준수량', '적용 필요예정수량'):
        assert config[column]['type_config']['type'] == 'number'
    assert not config["실제 발주수량"].get("disabled", False)
    assert config["추천 발주수량"]["disabled"] and config["계산 발주수량"]["disabled"]
    assert not at.selectbox
    first_id = proto.id
    updated = at.session_state['sims_export_tables']['widget'].copy()
    updated.loc[updated.index[0], '실제 발주수량'] = D(55)
    at.session_state['sims_export_tables'] = {'widget': updated}
    at.run()
    assert not at.exception and at.dataframe[0].proto.id == first_id
    many = pd.concat([updated] * 301, ignore_index=True)
    many['제품코드'] = [f'{i:05d}' for i in range(301)]
    at.session_state['sims_export_tables'] = {'widget': many}
    at.run()
    assert not at.exception and len(at.selectbox) == 1
    assert list(at.selectbox[0].options) == ['1', '2']


def test_panel_download_render_policy():
    import ast
    from pathlib import Path
    tree = ast.parse(Path('app/Lmstudio_SSAI_chat_main.py').read_text(encoding='utf-8-sig'))
    branch = next(n for n in ast.walk(tree) if isinstance(n, ast.If) and
        'reason == \'download_prepare\'' in ast.unparse(n.test) and 'selected_now' in ast.unparse(n.test))
    expression = compile(ast.Expression(branch.test), 'panel-policy', 'eval')
    for reason, action, blocked in [('download_prepare', '발주 계산', False), ('download_prepare', '제품정보 조회', True), ('chat_room_change', '발주 계산', True), ('order_edit', '발주 계산', False)]:
        assert eval(expression, {'reason': reason, 'selected_now': {'action': action}}) == blocked


def test_empty_parameter_bridge_and_check_mode():
    import app.services.rddbc_io_common as common
    with patch.object(common, "_resolve_query_func") as resolve:
        fn = resolve.return_value
        common.query_to_df("SELECT 1")
        fn.assert_called_once_with("SELECT 1", ())
    params, sources = fixture()
    params["query_mode"] = "확인 필요"
    frame = assemble_result(params, sources)
    assert frame.empty or (frame["계산상태"].str.contains("확인") | frame["발주단가"].isna() | frame["발주처코드"].eq("")).all()


def test_snapshot_scope_to_demand_scope():
    from types import SimpleNamespace
    from app.services.order_calculation_service import build_demand_params
    from app.services.analytics_sales_trend_service import _build_monthly_filters
    q = build_demand_params({"maker_nm": "삼진제약", "product_class_list": ["0031:01"],
        "product_di_list": ["0004:01"], "product_group_list": ["0013:01"]},
        SimpleNamespace(io_gu_codes=("001", "500", "601")))
    assert q["io_gu_list"] == ["500"] and q["product_ven_nm"] == "삼진제약"
    assert "product_class_list" not in q and q["dashboard_product_class_list"] == ["0031:01"]
    sql = _build_monthly_filters(q, {"prefix": "Rd21", "alias": "M", "out_prefixes": "'5','6'"})
    assert "Rd04_Physic_Tax" in sql and "Rd04_Physic_Gu" not in sql


def test_code_lookup_company_cache():
    import app.services.rddbc_io_common as common
    common._load_rddbc010_lookup_for_company.cache_clear()
    def lookup(*args):
        return pd.DataFrame({"Gcode": ["0013"], "Tcode": ["01"], "CodeName": [str(company.return_value)]})
    with patch("app.db.mssql_client.get_current_company_id", return_value=7) as company, \
         patch.object(common, "query_to_df", side_effect=lookup) as query:
        assert common._load_rddbc010_lookup()[1][("0013", "01")] == "7"
        company.return_value = 4
        assert common._load_rddbc010_lookup()[1][("0013", "01")] == "4"
        company.return_value = 7
        assert common._load_rddbc010_lookup()[1][("0013", "01")] == "7"
        assert query.call_count == 2
    common._load_rddbc010_lookup_for_company.cache_clear()


if __name__ == "__main__":
    tests = (test_horizon, test_monthly_allocation, test_quantities, test_demand_trend_adjustment, test_trend_applies_before_stock_pending_and_unit, test_assembly_edit_export, test_company_isolation, test_purchase_vendor_count_staff_and_sensitive_projection, test_routes_menu, test_panel_submission, test_price_conflict_and_field_status, test_scoped_timeout_and_no_retry, test_production_nlq_dispatch, test_editor_callback_ownership_and_cache)
    tests += (test_empty_parameter_bridge_and_check_mode, test_snapshot_scope_to_demand_scope)
    tests += (test_code_lookup_company_cache,)
    tests += (test_excel_numeric_round_trip, test_production_editor_render_boundary)
    tests += (test_streamlit_editor_widget,)
    tests += (test_panel_download_render_policy,)
    tests += (test_full_month_forecast_and_blank_pending,)
    tests += (test_horizon_settlement_boundary,)
    def test_readability_projection():
        params, sources = fixture()
        with patch('app.services.order_calculation_service.get_current_company_id', return_value=7):
            result = get_order_calculation_result(params, source_loader=lambda q: sources)
        full, display = result['df'], result['df_display']
        common = {'안전재고일수', '적정재고일수', '결제일자', '해당 월 전체 영업일수', '적용 필요 영업일수'}
        assert common.issubset(full.columns) and not common.intersection(display.columns)
        sequence = ['계산 발주수량', '추천 발주수량', '실제 발주수량', '재고수량', '입고예정수량']
        offset = list(display).index(sequence[0])
        assert list(display)[offset:offset + 5] == sequence
        assert '{' not in result['meta']['query_summary']
        assert list(display).index('추세') + 1 == list(display).index('계산 발주수량')
        assert list(display).count('추세') == 1
        assert result['meta']['elapsed_ms'] >= 0
        assert result['meta']['calculation_basis'].startswith('계산기준:')
        assert '계산기준:' in result['meta']['summary_md']
    tests += (test_readability_projection,)
    def test_order_unit_repetition():
        from app.services.order_calculation_contract import infer_order_unit
        def history(values):
            return [{'발주일자': f'202609{i+1:02d}', '발주수량': q} for i, q in enumerate(values)]
        assert infer_order_unit(history([10, 20, 10, 30, 20]))[0] == 10
        assert infer_order_unit(history([3, 7, 12, 20]))[0] is None
        assert infer_order_unit(history([10, 20]))[0] is None
        assert infer_order_unit(history([10, 20, 30]))[0] is None
        assert infer_order_unit(history(['1.5', 3, '1.5']))[0] is None
        assert infer_order_unit([])[0] is None
        assert recommend_quantity(D('23.4'), D(10), increasing=True) == 30
        assert recommend_quantity(D('23.4'), D(10), increasing=False) == 20
        assert recommend_quantity(D('7.8'), None, increasing=True) == 8
        assert recommend_quantity(D('7.8'), None, increasing=False) == 7
        params, sources = fixture()
        rows = history([10, 20, 10, 30])
        sources['order_history'] = pd.DataFrame([{**r, '제품코드': '00001', '발주거래처코드': '00100'} for r in rows])
        sources['order_history_complete'] = True
        row = assemble_result(params, sources).iloc[0]
        assert row['발주단위'] == 10 and row['계산 발주수량'] == 32 and row['추천 발주수량'] == 30
        updated = apply_actual_edits(pd.DataFrame([row]), pd.DataFrame([row]), {0: {'실제 발주수량': '40'}}).iloc[0]
        assert updated['계산 발주수량'] == 32 and updated['추천 발주수량'] == 30
        assert updated['실제 발주수량'] == 40 and updated['발주금액(부가세포함)'] == 4400
        sources['base']['출고빈도등급'] = 'A'
        sources['demand']['최근3개월평균수요수량'] = D(120)
        sources['demand']['직전3개월평균수요수량'] = D(100)
        sources['demand']['완료월수'] = 6
        assert assemble_result(params, sources).iloc[0]['추천 발주수량'] == 40
        sources['demand']['최근3개월평균수요수량'] = D(100)
        sources['order_history']['발주수량'] = 100
        zero = assemble_result(params, sources).iloc[0]
        assert zero['계산 발주수량'] == 32 and zero['추천 발주수량'] == 100
        assert '최소 발주단위' in zero['조달주의']
        assert assemble_result({**params, 'only_needed': True}, sources).iloc[0]['추천 발주수량'] == 100
    tests += (test_order_unit_repetition,)
    def test_order_response_timing():
        from app.ui import chat_middleware as chat
        payload = {'meta': {'order_calculation_editable': True,
            'request_started_at': '2026-09-14T11:00:00', 'request_started_monotonic': 100.0}}
        with patch.object(chat.time, 'monotonic', return_value=118.7):
            chat._attach_sims_response_timing(payload, {})
        assert payload['meta']['elapsed_ms'] == 18700
        assert 'request_started_monotonic' not in payload['meta']
        saved = dict(payload['meta'])
        chat._attach_sims_response_timing(payload, {})
        assert payload['meta'] == saved
    tests += (test_order_response_timing,)
    def test_editor_shared_fast_display():
        from types import SimpleNamespace
        from app.ui import order_calculation_editor as ui
        from app.ui.sims_table_display import build_sims_table_display_config
        params, sources = fixture()
        with patch('app.services.order_calculation_service.get_current_company_id', return_value=7):
            result = get_order_calculation_result(params, source_loader=lambda q: sources)
        full = result['df'].copy()
        full['발주단가'] = None
        full['발주금액(부가세포함)'] = float('nan')
        full['제약사'] = pd.NaT
        for column in ('월 기준 예상수량', '안전재고 기준수량', '기준 1영업일 예상수량'):
            full[column] = D('12.34567')
        display = full[result['df_display'].columns]
        expected, config, _, height = build_sims_table_display_config(
            display, action_name='발주 계산', meta=result['meta'], add_row_no=False, native_numeric_cells=True)
        saved = full.copy(deep=True)
        state = {'sims_export_tables': {'shared': full}, 'sims_tables': {'shared': display},
                 '__sims_current_table_source_key': 'shared'}
        captured = {}
        fake = SimpleNamespace(session_state=state, column_config=ui.st.column_config,
            data_editor=lambda df, **kw: captured.update(df=df, **kw), error=lambda *a: None)
        with patch.object(ui, 'st', fake), patch('app.db.mssql_client.get_current_company_id', return_value=7):
            assert ui.render_actual_quantity_editor({}, {**result['meta'], 'table_key': 'shared'}, download_uid='shared')
        for col in expected:
            if col != '실제 발주수량':
                pd.testing.assert_series_equal(captured['df'][col], expected[col])
                if col == '계산 발주수량' and config[col].get('type_config', {}).get('type') == 'number':
                    config[col]['type_config']['format'] = '%.2f'
                assert captured['column_config'][col] == config[col]
        for col in ('발주단가', '발주금액(부가세포함)'):
            assert captured['df'][col].isna().all()
            assert pd.api.types.is_numeric_dtype(captured['df'][col])
            assert captured['column_config'][col]['alignment'] == 'right'
        assert captured['column_config']['실제 발주수량']['alignment'] == 'right'
        for col in ('월 기준 예상수량', '안전재고 기준수량', '기준 1영업일 예상수량'):
            assert pd.api.types.is_numeric_dtype(captured['df'][col])
            assert captured['column_config'][col]['alignment'] == 'right'
        assert captured['placeholder'] == ' '
        assert captured['height'] == height
        pd.testing.assert_frame_equal(full, saved)
    tests += (test_editor_shared_fast_display,)
    def test_compact_order_header_and_initial_values():
        from app.ui import chat_middleware as chat
        from app.ui.sims_table_display import resolve_sims_numeric_display_kind
        params, sources = fixture()
        with patch('app.services.order_calculation_service.get_current_company_id', return_value=7):
            result = get_order_calculation_result(params, source_loader=lambda q: sources)
        view = chat._build_sims_result_header_view(result, result['meta'], result['df_display'])
        assert view['line1'] == view['line2'] == ''
        assert len(view['compact_lines']) == 2
        assert len(view['order_cards']) == 5
        assert [c[0] for c in view['order_condition_cards']] == ['발주일', '조회구분', '안전재고', '적정재고', '결제일']
        assert view['order_condition_cards'][0][1] == result['params']['order_date']
        assert view['compact_lines'][0] == ''
        assert view['order_cards'][4][1] == result['meta']['elapsed_ms'] / 1000
        assert any('월 영업일' in line for line in view['details'])
        with patch.object(chat.st, 'caption'), patch.object(chat.st, 'container') as container, patch.object(chat.st, 'expander'), patch.object(chat.st, 'markdown'), patch.object(chat, '_chat_metric_card') as card:
            chat._render_sims_result_header_view(view)
            assert card.call_count == 10
            assert container.call_args_list[0].kwargs == {'horizontal': True, 'gap': 'small'}
        assert resolve_sims_numeric_display_kind('계산 발주수량') == 'decimal2'
        for raw, unit, expected in ((D('6.56'), D(10), 10), (D('12.365'), D(100), 100)):
            row = calculate_quantities(stock=D(0), pending=D(0), safety_demand=D(1), horizon_demand=raw, unit=unit)
            assert row['추천 발주수량'] == row['실제 발주수량'] == expected
        unavailable = calculate_quantities(stock=D(0), pending=D(0), safety_demand=None, horizon_demand=None)
        assert unavailable['추천 발주수량'] is None and unavailable['실제 발주수량'] is None
    tests += (test_compact_order_header_and_initial_values,)
    def test_export_order_and_application_defaults():
        import inspect
        from app.services import order_calculation_service as order_service
        from app.services.order_calculation_service import apply_application_defaults, build_contract_price_params
        order_query = {'order_date': '2026-09-14', 'cost_apply_cd': '50002',
                       'stock_apply_cd': '50001', 'physic_cd': '83315',
                       'date_from': '20260815', 'date_to': '20260914',
                       'month_from': '202608', 'month_to': '202609'}
        contract_query = build_contract_price_params(order_query)
        assert not any(k in contract_query for k in ('date_from', 'date_to', 'month_from', 'month_to'))
        assert contract_query['ven_cd'] == '50002' and contract_query['as_of'] == '20260914'
        assert contract_query['physic_cd'] == '83315' and contract_query['stock_apply_cd'] == '50001'
        assert order_query['date_from'] == '20260815'
        assert apply_application_defaults({}) == {
            'cost_apply_cd': '50002', 'stock_apply_cd': '50001',
            'query_mode': '발주해당자료만', 'only_needed': True,
        }
        explicit = apply_application_defaults({'cost_apply_cd': '12345', 'stock_apply_nm': '지정처'})
        assert explicit['cost_apply_cd'] == '12345' and 'stock_apply_cd' not in explicit
        params, sources = fixture()
        with patch('app.services.order_calculation_service.get_current_company_id', return_value=7):
            result = get_order_calculation_result(params, source_loader=lambda q: sources)
        primary_front = ['제품코드', '제품명', '추세', '계산 발주수량', '추천 발주수량', '실제 발주수량',
                '재고수량', '입고예정수량', '발주처', '발주담당자', '제약담당자', '매입거래처수',
                '발주단가', '발주금액(부가세포함)', '월 기준 예상수량']
        assert list(result['df'].columns[:len(primary_front)]) == primary_front
        assert list(result['df_display'].columns[:len(primary_front)]) == primary_front
        assert '단가출처' in result['df'] and '발주단위 근거' in result['df']
        assert {'품목손익등급', '품목기여등급'}.issubset(result['df'].columns)
        assert {'품목손익등급', '품목기여등급'}.issubset(result['df_display'].columns)
        assert not {'손익등급', '기여도등급', '손익기여도'}.intersection(result['df'].columns)
        assert result['params']['cost_apply_cd'] == '50002'
        assert result['params']['stock_apply_cd'] == '50001'
        assert 'search_vendors_full' not in inspect.getsource(order_service.load_sources)
        assert "demand_params['order_stock_apply_cd']" not in inspect.getsource(order_service.load_sources)
    tests += (test_export_order_and_application_defaults,)
    def test_query_mode_is_display_only():
        params, sources = fixture()
        seen = []
        def loader(query):
            seen.append(dict(query))
            assert 'query_mode' not in query and 'only_needed' not in query
            return sources
        with patch('app.services.order_calculation_service.get_current_company_id', return_value=7):
            full = get_order_calculation_result({**params, 'query_mode': '전체', 'only_needed': True}, source_loader=loader)
            needed = get_order_calculation_result({**params, 'query_mode': '발주해당자료만', 'only_needed': False}, source_loader=loader)
        expected = full['df'].loc[full['df']['계산상태'].eq('발주해당')].reset_index(drop=True)
        pd.testing.assert_frame_equal(expected, needed['df'])
        assert full['params']['only_needed'] is False and needed['params']['only_needed'] is True
        assert seen[0] == seen[1]
    tests += (test_query_mode_is_display_only,)
    def test_panel_mode_reuses_company_source():
        from streamlit.testing.v1 import AppTest
        params, sources = fixture()
        calls = []
        def loader(query):
            calls.append(dict(query))
            query['cost_apply_nm'] = '단가 적용 거래처'
            query['stock_apply_nm'] = '재고 적용 거래처'
            return sources
        script = ('from app.sims.views.order_calculation_view import view_order_calculation\n'
                  "view_order_calculation({'order_date': '2026-09-14'})")
        with patch('app.db.mssql_client.get_current_company_id', return_value=7), \
             patch('app.services.order_calculation_service.get_current_company_id', return_value=7), \
             patch('app.sims.views.order_calculation_view.load_sources', side_effect=loader), \
             patch('app.sims.views.order_calculation_view.render_product_master_filters', return_value={}):
            app = AppTest.from_string(script).run()
            assert not app.exception
            assert [t.label for t in app.text_input[:4]] == ['단가적용처코드', '단가적용처명', '재고적용처코드', '재고적용처명']
            app.button[0].click().run()
            assert not app.exception and len(calls) == 1
            app.selectbox[0].select('전체')
            app.button[0].click().run()
            assert not app.exception and len(calls) == 1
            assert app.text_input[1].value == '단가 적용 거래처'
            assert app.text_input[3].value == '재고 적용 거래처'
            app.button[0].click().run()
            assert not app.exception and len(calls) == 2
    tests += (test_panel_mode_reuses_company_source,)
    for test in tests:
        test()
        print("PASS", test.__name__)
    print(f"PASS order calculation {len(tests)}/{len(tests)} offline")
